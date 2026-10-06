"""Read and validate an entire roster before any login account is created."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from app.models import User, Student, Faculty, Course, Curriculum

MAX_FILE_BYTES = 1024 * 1024
MAX_RECORDS = 200
BASE_HEADERS = {
    'students': ['full_name', 'username', 'email', 'enrollment_number', 'programme', 'pattern', 'semester', 'admission_year'],
    'staff': ['full_name', 'username', 'email', 'employee_id', 'department', 'qualification'],
}
HEADERS = {
    'students': BASE_HEADERS['students'] + ['mobile_number', 'gender', 'date_of_birth', 'address', 'practical_batch'],
    'staff': BASE_HEADERS['staff'] + ['mobile_number', 'gender', 'joining_date'],
}


def _read_excel(raw):
    """Bound ZIP expansion and cell counts before reading untrusted Excel data."""
    from openpyxl import load_workbook
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            files = archive.infolist()
            if len(files) > 1000 or sum(item.file_size for item in files) > 8 * 1024 * 1024:
                raise ValueError('The Excel workbook is too large. Use the blank PRAVA template.')
            if any('vbaproject' in item.filename.lower() or 'externallinks/' in item.filename.lower() for item in files):
                raise ValueError('Remove macros and external workbook links before uploading.')
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook['Roster'] if 'Roster' in workbook.sheetnames else workbook.worksheets[0]
            if (sheet.max_row or 0) > MAX_RECORDS + 1 or (sheet.max_column or 0) > max(map(len, HEADERS.values())):
                raise ValueError('The Roster sheet must contain a header and at most 200 data rows, using only template columns.')
            values = []
            for index, row in enumerate(sheet.iter_rows(), 1):
                if index > MAX_RECORDS + 1 or len(row) > max(map(len, HEADERS.values())):
                    raise ValueError('The Roster sheet must contain a header and at most 200 data rows, using only template columns.')
                result = []
                for cell in row:
                    if cell.data_type in {'f', 'e'}:
                        raise ValueError(f'Cell {cell.coordinate}: paste values instead of formulas or Excel errors.')
                    value = cell.value
                    if isinstance(value, datetime):
                        value = value.date().isoformat()
                    elif isinstance(value, date):
                        value = value.isoformat()
                    elif isinstance(value, float) and value.is_integer():
                        value = int(value)
                    result.append('' if value is None else str(value))
                values.append(result)
            return values
        finally:
            workbook.close()
    except ValueError:
        raise
    except (BadZipFile, KeyError, IndexError, OSError, TypeError) as exc:
        raise ValueError('Upload a valid .xlsx workbook using the PRAVA template.') from exc
    except Exception as exc:
        # XML parser failures must remain a validation response, never a partial import.
        raise ValueError('The Excel workbook could not be read. Save it as a new .xlsx file and try again.') from exc


def validate_roster(raw, kind, filename='roster.csv'):
    if kind not in HEADERS:
        raise ValueError('Choose Students or Staff.')
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError('The file must be smaller than 1 MB.')
    extension = Path(filename or '').suffix.lower()
    if extension == '.xlsx':
        values = _read_excel(raw)
    elif extension == '.csv':
        try:
            values = list(csv.reader(io.StringIO(raw.decode('utf-8-sig')), strict=True))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError('Upload a valid UTF-8 CSV file.') from exc
    else:
        raise ValueError('Upload an Excel .xlsx file or a UTF-8 .csv file.')
    if not values or values[0] not in (HEADERS[kind], BASE_HEADERS[kind]):
        raise ValueError('Use the exact columns and order in the downloadable template.')
    headers = values[0]
    numbered_rows = [(number, cells) for number, cells in enumerate(values[1:], 2) if any(str(value).strip() for value in cells)]
    if not 1 <= len(numbered_rows) <= MAX_RECORDS:
        raise ValueError('Each import must contain between 1 and 200 records.')
    rows = []
    for number, cells in numbered_rows:
        if len(cells) != len(headers):
            raise ValueError(f'Row {number}: The number of columns does not match the template.')
        rows.append((number, {**dict.fromkeys(HEADERS[kind], ''), **dict(zip(headers, cells))}))
    return _validate_rows(rows, kind)


def _unique(base, used, limit):
    candidate, suffix = base[:limit], 1
    while candidate.casefold() in used:
        suffix += 1
        ending = f'-{suffix}'
        candidate = base[:limit - len(ending)] + ending
    return candidate


def _date_value(row, field, prefix):
    if not row.get(field):
        return
    try:
        parsed = date.fromisoformat(row[field])
    except ValueError as exc:
        raise ValueError(prefix + f'{field} must use YYYY-MM-DD.') from exc
    if parsed > date.today():
        raise ValueError(prefix + f'{field} cannot be in the future.')
    row[field] = parsed.isoformat()


def _validate_rows(numbered_rows, kind):
    identity_key = 'enrollment_number' if kind == 'students' else 'employee_id'
    model = Student if kind == 'students' else Faculty
    # Read once for an entire file, including 160-student registrations.
    existing_names = {value.casefold() for (value,) in User.query.with_entities(User.username)}
    existing_emails = {value.casefold() for (value,) in User.query.with_entities(User.email)}
    existing_ids = {value.casefold() for (value,) in model.query.with_entities(getattr(model, identity_key))}
    seen = {'username': set(), 'email': set(), 'identity': set()}
    reserved_names = {str(row.get('username', '')).strip().casefold() for _, row in numbered_rows if row.get('username')}
    reserved_ids = {str(row.get(identity_key, '')).strip().casefold() for _, row in numbered_rows if row.get(identity_key)}
    courses = {course.code: course for course in Course.query.filter_by(is_active=True)}
    curricula = {(item.course_id, item.pattern): item for item in Curriculum.query.all()}
    result = []
    for number, source in numbered_rows:
        prefix = f'Row {number}: '
        row = {key: str(source.get(key, '')).strip() for key in HEADERS[kind]}
        required = set(BASE_HEADERS[kind]) - {'username', identity_key, 'qualification'}
        for key in required:
            if not row[key]:
                raise ValueError(prefix + f'{key} is required.')
        row['email'] = row['email'].lower()
        if len(row['full_name']) > 120 or len(row['email']) > 120 or not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', row['email']):
            raise ValueError(prefix + 'Check the name and email address.')
        if any('\x00' in value or value.startswith('=') for value in row.values()):
            raise ValueError(prefix + 'Use plain values, without formulas or control characters.')
        row['username'] = row['username'].lower()
        if not row['username']:
            slug = unicodedata.normalize('NFKD', row['full_name']).encode('ascii', 'ignore').decode().lower()
            slug = re.sub(r'[^a-z0-9]+', '.', slug).strip('.')
            base = ('student.' if kind == 'students' else 'staff.') + (slug or 'account')
            row['username'] = _unique(base, existing_names | reserved_names | seen['username'], 80)
        if not re.fullmatch(r'[a-z0-9_.-]{3,80}', row['username']):
            raise ValueError(prefix + 'Username must contain 3–80 letters, digits, dots, hyphens or underscores.')
        # Match the manual student/faculty editors before preview and persistence.
        row[identity_key] = row[identity_key].upper()
        if not row[identity_key]:
            base = ('PRV-STU-' if kind == 'students' else 'PRV-STF-') + str(date.today().year)
            row[identity_key] = _unique(base, existing_ids | reserved_ids | seen['identity'], 50)
        identity = row[identity_key]
        if len(identity) > 50:
            raise ValueError(prefix + 'Enrollment / employee number is too long.')
        for key, value in [('username', row['username']), ('email', row['email']), ('identity', identity.casefold())]:
            if value in seen[key]:
                raise ValueError(prefix + f'Duplicate {key} in this file.')
            seen[key].add(value)
        if row['username'] in existing_names or row['email'] in existing_emails:
            raise ValueError(prefix + 'This username or email already exists. Edit the existing account.')
        if identity.casefold() in existing_ids:
            raise ValueError(prefix + 'This enrollment / employee number already exists.')
        if row['mobile_number'] and not re.fullmatch(r'\+?[0-9]{10,15}', row['mobile_number']):
            raise ValueError(prefix + 'Mobile number must contain 10–15 digits, optionally starting with +.')
        if row['gender']:
            row['gender'] = row['gender'].title()
            if row['gender'] not in {'Female', 'Male', 'Other'}:
                raise ValueError(prefix + 'Gender must be Female, Male or Other, or left blank.')
        if kind == 'students':
            course = courses.get(row['programme'].upper())
            if not course or (course.id, row['pattern']) not in curricula:
                raise ValueError(prefix + 'Programme and pattern must match a configured curriculum.')
            try:
                semester, year = int(row['semester']), int(row['admission_year'])
                if not (1 <= semester <= course.total_semesters and 1997 <= year <= date.today().year + 1):
                    raise ValueError()
            except ValueError as exc:
                raise ValueError(prefix + 'Check the semester and admission year.') from exc
            row['programme'] = course.code
            row['practical_batch'] = row['practical_batch'].upper()
            if row['practical_batch'] and not re.fullmatch(r'[A-Z0-9-]{1,10}', row['practical_batch']):
                raise ValueError(prefix + 'Practical batch must contain up to 10 letters or digits.')
            if len(row['address']) > 1000:
                raise ValueError(prefix + 'Address must be at most 1000 characters.')
            _date_value(row, 'date_of_birth', prefix)
        else:
            if len(row['department']) > 120 or len(row['qualification']) > 120:
                raise ValueError(prefix + 'Department and qualification must be at most 120 characters.')
            _date_value(row, 'joining_date', prefix)
        result.append(row)
    return result


def encode_roster(rows, kind):
    """Internal lossless serialization, also used for blank CSV templates."""
    out = io.StringIO(newline='')
    writer = csv.DictWriter(out, fieldnames=HEADERS[kind])
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode('utf-8')
